"""BOX step 15 — build the owner's edit deck: one slide per bench with its visual-check map (read-only, no DB).

    python -m scripts.box_extent_pptx --out "<dir>/extents_2026-10-16.pptx" BS2_C=docs/box/step15-BS2C-2026-10-16/extents ...
    python -m scripts.box_extent_pptx --out <deck> --basin Midland WCA_1=<20a>/extents \
        "WCB_2=<20a>/extents@<20a>/wcb_split/visual_WCB_2_clipped.png" --note "WCB_2=D136: clipped at the WCB split line, 2,041 sq mi"

Each slide: the PNG at the exact placement Michael used for bs3_extents.pptx (left 2,667,000 EMU, 6,858,000 EMU square,
uncropped — scripts.box_family_extent_edits maps freeforms back through the PNG's recorded pixel frame) and a text box
with the numbers. Edit convention of record (step 15, 2026-10-16): the owner's red CLOSED loops ARE the extent, the black
generated outline is ignored; a slide left untouched accepts the generated extent. ``BENCH=<dir>@<png>`` uses another map
(its <png stem>.frame.json sits beside it, e.g. the D136 clipped WCB_2); ``--note BENCH=text`` replaces the area line. Writes <out>.manifest.json (slide -> bench, PNG md5) for the ingest. Needs python-pptx (not
in the shared venv: run with PYTHONPATH pointing at a private install).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

PIC_LEFT, PIC_SIDE = 2_667_000, 6_858_000
SLIDE_W, SLIDE_H = 12_192_000, 6_858_000


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("items", nargs="+", help="BENCH=<extents dir>[@<png>] (visual_<BENCH>.png + summary.json there)")
    ap.add_argument("--note", action="append", default=[], help="BENCH=text: replaces the slide's generated-area line")
    ap.add_argument("--basin", default="", help="label on each slide title (e.g. Midland)")
    a = ap.parse_args(argv)
    from pptx import Presentation
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    manifest = {"slides": {}}
    for k, item in enumerate(a.items, start=1):
        bench, d = item.split("=", 1)
        d, _, over = d.partition("@")
        d = Path(d)
        png = Path(over) if over else d / f"visual_{bench}.png"
        notes = dict(n.split("=", 1) for n in a.note)
        st = json.loads((d / "summary.json").read_text(encoding="utf-8"))["pools"][bench]
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), Emu(PIC_LEFT), Emu(0), Emu(PIC_SIDE), Emu(PIC_SIDE))
        tb = s.shapes.add_textbox(Emu(80_000), Emu(300_000), Emu(PIC_LEFT - 160_000), Emu(5_000_000))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = [
            (f"{a.basin} {bench}".strip(), 20, True),
            (notes.get(bench, f"generated one-body D27 extent: {st['extent_sqmi']:,.0f} sq mi"), 11, False),
            (f">= 2016 laterals: {st['n_evidence']:,}; inside {st.get('evidence_inside', 0):,}", 11, False),
            (f"bodies >= 50 laterals: {st.get('n_bodies_ge_min', 0)} (dashed; body 1 = the D27 body)", 11, False),
            (f"isolated tests out: {st['n_isolated_tests']:,}", 11, False),
            ("", 8, False),
            ("Edits (pen or freeform, red):", 11, True),
            ("your closed loops ARE the extent", 11, False),
            ("(the black outline is ignored)", 11, False),
            ("no drawing = accept as generated", 11, False),
        ]
        for i, (txt, size, bold) in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = txt
            para.font.size = Pt(size)
            para.font.bold = bold
        manifest["slides"][str(k)] = {"bench": bench, "extents_dir": str(d), "png": png.name, "png_path": str(png), "png_md5": hashlib.md5(png.read_bytes()).hexdigest()}
    a.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(a.out))
    Path(str(a.out) + ".manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    print(f"wrote {a.out} ({len(a.items)} slides) + manifest")
    return 0


if __name__ == "__main__":
    sys.exit(main())
